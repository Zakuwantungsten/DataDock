from datetime import date
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple, Union

try:
    from .config import get_column_aliases
except ImportError:  # pragma: no cover - supports running as a script
    from config import get_column_aliases


class ExcelMerger:
    def __init__(self) -> None:
        self.loaded_files: List[Path] = []
        self.errors: List[str] = []

    def set_files(self, file_paths: Iterable[Path]) -> None:
        self.loaded_files = [Path(path) for path in file_paths]

    def _normalize_column_name(self, name: str) -> str:
        return str(name).strip().lower()

    def _resolve_columns(
        self,
        df_columns: Iterable[str],
        aliases_map: Dict[str, List[str]],
    ) -> Tuple[List[str], Dict[str, str], List[str]]:
        normalized = {
            self._normalize_column_name(column): column
            for column in df_columns
        }
        actual_columns: List[str] = []
        rename_map: Dict[str, str] = {}
        missing_targets: List[str] = []

        for target_column, aliases in aliases_map.items():
            candidates = [target_column] + list(aliases)
            matched_column = None
            for candidate in candidates:
                key = self._normalize_column_name(candidate)
                if key in normalized:
                    matched_column = normalized[key]
                    break

            if matched_column is None:
                missing_targets.append(target_column)
                continue

            actual_columns.append(matched_column)
            rename_map[matched_column] = target_column

        return actual_columns, rename_map, missing_targets

    def _get_last_sheet_name(self, file_path: Path) -> Optional[str]:
        try:
            workbook = load_workbook(file_path, read_only=True, data_only=True)
        except Exception:
            return None

        if not workbook.sheetnames:
            return None

        return workbook.sheetnames[-1]

    def _detect_header_row(
        self,
        file_path: Path,
        sheet_name: str,
        aliases_map: Dict[str, List[str]],
        max_rows: int = 10,
    ) -> Optional[int]:
        try:
            preview = pd.read_excel(
                file_path,
                sheet_name=sheet_name,
                header=None,
                nrows=max_rows,
                engine="openpyxl",
            )
        except Exception:
            return None

        alias_lookup: Dict[str, str] = {}
        for target, aliases in aliases_map.items():
            alias_lookup[self._normalize_column_name(target)] = target
            for alias in aliases:
                alias_lookup[self._normalize_column_name(alias)] = target

        best_row = None
        best_matches = 0
        for row_index in range(len(preview.index)):
            row_values = preview.iloc[row_index].tolist()
            matched_targets = set()
            for value in row_values:
                if value is None:
                    continue
                key = self._normalize_column_name(value)
                target = alias_lookup.get(key)
                if target:
                    matched_targets.add(target)

            if len(matched_targets) > best_matches:
                best_matches = len(matched_targets)
                best_row = row_index

        if best_matches == 0:
            return None

        return best_row

    def extract_columns(
        self,
        file_paths: Iterable[Path],
        aliases_map: Optional[Dict[str, List[str]]] = None,
        header_row: Optional[int] = None,
    ) -> Dict[str, pd.DataFrame]:
        tables: Dict[str, pd.DataFrame] = {}
        errors: List[str] = []
        if aliases_map is None:
            aliases_map = get_column_aliases()

        for path in file_paths:
            file_path = Path(path)
            if not file_path.exists():
                errors.append(f"{file_path.name}: file not found")
                continue

            resolved_sheet = self._get_last_sheet_name(file_path)
            if resolved_sheet is None:
                errors.append(f"{file_path.name}: no sheets found")
                continue

            resolved_header = header_row
            if resolved_header is None:
                resolved_header = self._detect_header_row(
                    file_path,
                    resolved_sheet,
                    aliases_map,
                )
            if resolved_header is None:
                errors.append(
                    f"{file_path.name}: header row not detected"
                )
                continue

            try:
                df = pd.read_excel(
                    file_path,
                    sheet_name=resolved_sheet,
                    header=resolved_header,
                    engine="openpyxl",
                )
            except Exception as exc:
                errors.append(f"{file_path.name}: {exc}")
                continue

            df.columns = df.columns.astype(str).str.strip()
            actual_columns, rename_map, missing_targets = self._resolve_columns(
                df.columns,
                aliases_map,
            )
            if missing_targets:
                errors.append(
                    f"{file_path.name}: missing aliases for {missing_targets}"
                )
                continue

            extracted = df.loc[:, actual_columns].rename(columns=rename_map)
            extracted = extracted[list(aliases_map.keys())]
            tables[file_path.name] = extracted

        self.errors = errors
        return tables

    def merge_files(
        self,
        aliases_map: Optional[Dict[str, List[str]]] = None,
        header_row: Optional[int] = None,
    ) -> Dict[str, pd.DataFrame]:
        return self.extract_columns(
            self.loaded_files,
            aliases_map,
            header_row=header_row,
        )

    def format_tables(
        self,
        tables: Dict[str, pd.DataFrame],
    ) -> Dict[str, pd.DataFrame]:
        formatted: Dict[str, pd.DataFrame] = {}
        required_columns = [
            "Truck",
            "Trailer",
            "Position",
            "Status",
            "Departure Date",
        ]
        output_columns = [
            "SN",
            "Truck",
            "Trailer",
            "Position",
            "Status",
            "Type",
            "Return",
            "DSJ",
        ]

        for name, table in tables.items():
            missing = [col for col in required_columns if col not in table.columns]
            if missing:
                self.errors.append(
                    f"{name}: missing required columns {missing}"
                )
                continue

            table = table.copy()
            table[required_columns] = table[required_columns].replace(
                r"^\s*$", pd.NA, regex=True
            )
            table = table.dropna(how="all", subset=required_columns).reset_index(drop=True)

            departure_dates = pd.to_datetime(
                table["Departure Date"],
                errors="coerce",
            )
            invalid_mask = departure_dates.isna()
            if invalid_mask.any():
                self.errors.append(
                    f"{name}: {int(invalid_mask.sum())} rows missing/invalid Departure Date"
                )

            today = pd.Timestamp.now().normalize()
            dsj = (today - departure_dates).dt.days
            dsj = dsj.where(~invalid_mask, other=0).astype("Int64")

            formatted_table = pd.DataFrame(
                {
                    "SN": range(1, len(table) + 1),
                    "Truck": table["Truck"],
                    "Trailer": table["Trailer"],
                    "Position": table["Position"],
                    "Status": table["Status"],
                    "Type": "Flatbed",
                    "Return": "NA",
                    "DSJ": dsj,
                }
            )
            formatted[name] = formatted_table[output_columns]

        return formatted

    def aggregate(
        self,
        data: Union[pd.DataFrame, Dict[str, pd.DataFrame]],
        group_by: Optional[List[str]] = None,
        agg_map: Optional[Dict[str, str]] = None,
    ) -> Union[pd.DataFrame, Dict[str, pd.DataFrame]]:
        if not group_by or not agg_map:
            return data

        if isinstance(data, dict):
            aggregated: Dict[str, pd.DataFrame] = {}
            for name, table in data.items():
                aggregated[name] = table.groupby(
                    group_by,
                    dropna=False,
                    as_index=False,
                ).agg(agg_map)
            return aggregated

        return data.groupby(group_by, dropna=False, as_index=False).agg(agg_map)

    def export_excel(self, df: pd.DataFrame, output_path: Path) -> Path:
        output_file = Path(output_path)
        output_file.parent.mkdir(parents=True, exist_ok=True)
        df.to_excel(output_file, index=False, engine="openpyxl")
        return output_file

    def export_excel_tables(
        self,
        tables: Dict[str, pd.DataFrame],
        output_path: Path,
        sheet_name: str = "Import Report",
        table_spacing: int = 2,
        report_title: str = "Import Report",
        report_date: Optional[str] = None,
    ) -> Path:
        output_file = Path(output_path)
        output_file.parent.mkdir(parents=True, exist_ok=True)

        max_columns = max((len(table.columns) for table in tables.values()), default=0)
        if report_date is None:
            report_date = date.today().strftime("%d-%b")

        start_row = 0
        report_header_row: Optional[int] = None
        table_ranges: List[Dict[str, object]] = []
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            if max_columns > 0:
                header_row = [""] * max_columns
                header_row[0] = report_title
                header_row[-1] = report_date
                pd.DataFrame([header_row]).to_excel(
                    writer,
                    sheet_name=sheet_name,
                    startrow=start_row,
                    index=False,
                    header=False,
                )
                report_header_row = start_row + 1
                start_row += 2

            for title, table in tables.items():
                display_title = Path(title).stem
                title_row = start_row
                pd.DataFrame([[display_title]]).to_excel(
                    writer,
                    sheet_name=sheet_name,
                    startrow=title_row,
                    index=False,
                    header=False,
                )
                title_row_excel = title_row + 1
                start_row = title_row + 1
                header_row = start_row
                table.to_excel(
                    writer,
                    sheet_name=sheet_name,
                    startrow=header_row,
                    index=False,
                )
                header_row_excel = header_row + 1
                table_ranges.append(
                    {
                        "title_row": title_row_excel,
                        "header_row": header_row_excel,
                        "row_count": len(table),
                        "col_count": len(table.columns),
                        "headers": list(table.columns),
                    }
                )
                start_row = header_row + len(table) + 1 + table_spacing

            ws = writer.sheets.get(sheet_name)
            if ws is not None and table_ranges:
                # Apply consistent styling to report tables.
                font_name = "Comic Sans MS"
                body_font = Font(name=font_name, color="000000")
                header_font = Font(name=font_name, bold=True, color="1F4E79")
                title_font = Font(name=font_name, bold=True, color="1F4E79")
                header_fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
                center = Alignment(horizontal="center", vertical="center", wrap_text=False)
                thin = Side(style="thin", color="9BA3AF")
                border = Border(left=thin, right=thin, top=thin, bottom=thin)

                target_headers = {"sn", "status", "position", "truck", "trailer", "type", "return", "dsj"}

                if report_header_row is not None:
                    if max_columns > 1:
                        ws.merge_cells(
                            start_row=report_header_row,
                            start_column=1,
                            end_row=report_header_row,
                            end_column=max_columns - 1,
                        )
                    for col in range(1, max_columns + 1):
                        cell = ws.cell(row=report_header_row, column=col)
                        cell.font = title_font
                        cell.alignment = center

                for table_info in table_ranges:
                    title_row_excel = int(table_info["title_row"])
                    header_row_excel = int(table_info["header_row"])
                    row_count = int(table_info["row_count"])
                    col_count = int(table_info["col_count"])
                    headers = [str(h) for h in table_info["headers"]]

                    if col_count > 1:
                        ws.merge_cells(
                            start_row=title_row_excel,
                            start_column=1,
                            end_row=title_row_excel,
                            end_column=col_count,
                        )
                    title_cell = ws.cell(row=title_row_excel, column=1)
                    title_cell.font = title_font
                    title_cell.alignment = center

                    end_row = header_row_excel + row_count
                    for row in range(header_row_excel, end_row + 1):
                        for col in range(1, col_count + 1):
                            cell = ws.cell(row=row, column=col)
                            cell.font = body_font
                            cell.alignment = center
                            cell.border = border

                    for col, header in enumerate(headers, start=1):
                        cell = ws.cell(row=header_row_excel, column=col)
                        cell.font = header_font
                        if header.strip().lower() in target_headers:
                            cell.fill = header_fill

                default_width = 10
                max_lengths: Dict[int, int] = {}
                for table in tables.values():
                    for col_index, header in enumerate(table.columns, start=1):
                        max_len = max_lengths.get(col_index, 0)
                        max_len = max(max_len, len(str(header)))
                        values = table.iloc[:, col_index - 1].tolist()
                        for value in values:
                            if value is None or (isinstance(value, float) and pd.isna(value)):
                                continue
                            if pd.isna(value):
                                continue
                            max_len = max(max_len, len(str(value)))
                        max_lengths[col_index] = max_len

                for col_index in range(1, max_columns + 1):
                    max_len = max_lengths.get(col_index, 0)
                    width = max_len + 2
                    if width > default_width:
                        column_letter = get_column_letter(col_index)
                        ws.column_dimensions[column_letter].width = width

        return output_file

    def export_pdf(
        self,
        df: pd.DataFrame,
        output_path: Path,
        title: Optional[str] = None,
    ) -> Path:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except Exception as exc:  # pragma: no cover - optional dependency
            raise ImportError(
                "PDF export requires 'reportlab'. Install it with: pip install reportlab"
            ) from exc

        output_file = Path(output_path)
        output_file.parent.mkdir(parents=True, exist_ok=True)

        data: List[List[str]] = [df.columns.astype(str).tolist()]
        data.extend(df.astype(str).values.tolist())

        doc = SimpleDocTemplate(str(output_file), pagesize=landscape(letter))
        styles = getSampleStyleSheet()
        elements = []

        if title:
            elements.append(Paragraph(title, styles["Title"]))
            elements.append(Spacer(1, 12))

        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ]
            )
        )
        elements.append(table)
        doc.build(elements)
        return output_file
    
